import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from shortlister import config, llm, store, vault  # noqa: E402

IDENTITY = {
    "name": "Priya Raghavan",
    "email": "priya.raghavan88@gmail.com",
    "phone": "+91 98765 43210",
    "colleges": ["IIT Bombay", "Narsee Monjee College of Commerce", "Jamnabai Narsee School"],
}


@pytest.fixture
def env(tmp_path, monkeypatch):
    """Isolated DB, vault, outbox and applications folder. Default env: no DRY_RUN set."""
    for var in ("DRY_RUN", "TEST_RECIPIENT", "TEST_WHATSAPP", "RESEND_API_KEY", "SCHEDULING_LINK", "DATABASE_URL",
                "KARGO_DEMO", "VAULT_KEY", "SLACK_WEBHOOK_URL", "HIRING_IMAP_HOST", "KARGO_LLM_PROVIDER", "KARGO_MODEL",
                "GEMINI_API_KEY", "GOOGLE_API_KEY", "GEMINI_PAID_TIER", "FROM_EMAIL", "REPLY_TO", "KARGO_STORAGE"):
        monkeypatch.delenv(var, raising=False)
    apps = tmp_path / "applications"
    apps.mkdir()
    monkeypatch.setattr(config, "APPLICATIONS_DIR", apps)
    monkeypatch.setattr(config, "META_DIR", apps / ".meta")
    monkeypatch.setattr(config, "OUTBOX_DIR", tmp_path / "outbox")
    monkeypatch.setattr(config, "VAULT_DIR", tmp_path / "vault")
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "backups")
    monkeypatch.setattr(vault, "VAULT_FILE", tmp_path / "vault" / "identities.json")
    store.set_db_path(tmp_path / "data" / "kargo.db")
    # Most tests check what happens once an email goes out; the undo window has its own tests.
    from shortlister import settings

    settings.put("undo_minutes", 0, "test")
    yield tmp_path
    store.set_db_path(config.DB_PATH)


def write_cv_docx(path: Path, role_line="Applying for: Product Manager", body=None):
    import docx

    d = docx.Document()
    d.add_paragraph(IDENTITY["name"])
    d.add_paragraph(f"{IDENTITY['email']}  |  {IDENTITY['phone']}  |  Andheri West, Mumbai")
    d.add_paragraph("Address: Flat 12, Sunshine CHS, SV Road, Andheri West, Mumbai 400058")
    d.add_paragraph("Date of Birth: 14/03/1994  |  Gender: Female  |  Marital Status: Married")
    d.add_paragraph(role_line)
    d.add_paragraph("EXPERIENCE")
    for line in body or [
        "Operations Executive · Seahawk Freight Forwarders, Nhava Sheva · Jun 2016 – May 2019",
        "Ran daily shipment documentation and carrier coordination on the port floor for 40 exporters.",
        "Built a WhatsApp-to-sheet tracker for container status that the other two branches adopted on their own.",
        "Product Manager · Tracklane (seed-stage logistics SaaS) · Jun 2021 – Present",
        "Sole PM reporting to the CEO; launched the carrier-booking module 0 to 1.",
        "When the ICEGATE outage hit in March 2023 she rerouted 60 shipping bills manually within 6 hours.",
        "Killed the rate-card feature after 8 weeks of low use and wrote the post-mortem that changed our discovery process.",
    ]:
        d.add_paragraph(line)
    d.add_paragraph("EDUCATION")
    d.add_paragraph("B.Tech, Mechanical Engineering  ·  IIT Bombay  ·  2011–2015")
    d.add_paragraph("MBA  ·  Narsee Monjee College of Commerce  ·  2019–2021")
    d.add_paragraph("Schooling: Jamnabai Narsee School, Mumbai")
    d.save(str(path))


def good_classification(payload: dict) -> dict:
    """A deterministic stand-in for the LLM: codes with quotes that exist in the CV."""
    q = {
        "P1": "Ran daily shipment documentation and carrier coordination on the port floor for 40 exporters.",
        "P2": "Built a WhatsApp-to-sheet tracker for container status that the other two branches adopted on their own.",
        "P3": "Sole PM reporting to the CEO; launched the carrier-booking module 0 to 1.",
        "P4": "rerouted 60 shipping bills manually within 6 hours.",
        "P5": "Killed the rate-card feature after 8 weeks of low use and wrote the post-mortem",
    }
    code = lambda c, e=None: {"code": c, "evidence": e}  # noqa: E731
    cid = payload["messages"][0]["content"].split("candidate_id: ")[1].split("\n")[0]
    return {
        "candidate_id": cid, "pm_years": 3.3, "pm_years_working": "fixture value",
        "city": "Mumbai", "open_to_relocate": "yes",
        "layer_a": {**{k: code(1, v) for k, v in q.items()}, "N1": code(0), "N2": code(0)},
        "layer_b": {
            "PM": {"zero_to_one_early_stage": code(1, q["P3"]), "shipped_with_adoption": code(0.5, q["P2"]),
                   "b2b_operational_users": code(1, q["P1"])},
            "SPM": {"integration_platform_data": code(0), "build_config_avoid_calls": code(0),
                    "early_stage_unstructured": code(1, q["P3"])},
        },
        "confidence": "high", "confidence_reason": "clear dates",
        "rationale": "Hands-on port operations before product. Owned the carrier-booking module as sole PM.",
        "probes": ["a?", "b?", "c?"],
    }


class FakeClassifier:
    """Records every payload that would have gone to the API."""

    def __init__(self, fn=good_classification):
        self.fn = fn
        self.payloads = []

    def __call__(self, payload):
        self.payloads.append(payload)
        obj = self.fn(payload)
        assert not llm.validate(obj), llm.validate(obj)
        return json.loads(json.dumps(obj)), None
